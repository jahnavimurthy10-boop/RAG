import asyncio
from io import BytesIO

import pytest
from docx import Document as DocxDocument

from app.services.documents import extract_text, split_into_chunks
from app.services.rag import RagService, cosine_similarity


def test_split_into_chunks_respects_size_and_overlaps() -> None:
    chunks = split_into_chunks("abcdefghij\nklmnopqrst", chunk_size=12, overlap=3)

    assert all(len(chunk) <= 12 for chunk in chunks)
    assert chunks[0] == "abcdefghij"
    assert chunks[1].startswith("hij")


def test_split_into_chunks_rejects_invalid_overlap() -> None:
    with pytest.raises(ValueError, match="overlap"):
        split_into_chunks("some text", chunk_size=10, overlap=10)


def test_cosine_similarity_ranks_direction_and_zero_vectors() -> None:
    assert cosine_similarity([1, 0], [4, 0]) == pytest.approx(1.0)
    assert cosine_similarity([1, 0], [0, 1]) == pytest.approx(0.0)
    assert cosine_similarity([0, 0], [1, 2]) == 0.0


def test_cosine_similarity_rejects_mismatched_dimensions() -> None:
    with pytest.raises(ValueError, match="dimensions"):
        cosine_similarity([1], [1, 2])


def test_extract_text_reads_docx_paragraphs_and_tables() -> None:
    document = DocxDocument()
    document.add_paragraph("Workshop notes")
    row = document.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text = "Retrieval"
    row.cells[1].text = "Generation"
    content = BytesIO()
    document.save(content)

    text = extract_text("notes.docx", content.getvalue())

    assert "Workshop notes" in text
    assert "Retrieval | Generation" in text


class FakeOllama:
    def __init__(self) -> None:
        self.prompt = ""

    async def embed(self, texts: list[str]) -> list[list[float]]:
        vectors = []
        for text in texts:
            normalized = text.lower()
            vectors.append(
                [1.0, 0.0] if "volcano" in normalized else [0.0, 1.0]
            )
        return vectors

    async def answer(self, prompt: str) -> str:
        self.prompt = prompt
        return "The notes describe a volcano."


def test_rag_retrieves_only_selected_document_context() -> None:
    async def run() -> None:
        client = FakeOllama()
        service = RagService(client)
        volcano = await service.add_document("volcano.txt", "A volcano can erupt.")
        await service.add_document("ocean.txt", "The ocean contains salt water.")

        answer, sources = await service.answer("Tell me about the volcano.", [volcano["id"]])

        assert answer == "The notes describe a volcano."
        assert [source["document"] for source in sources] == ["volcano.txt"]
        assert "A volcano can erupt." in client.prompt
        assert "salt water" not in client.prompt

    asyncio.run(run())


def test_rag_does_not_treat_empty_selection_as_all_documents() -> None:
    async def run() -> None:
        service = RagService(FakeOllama())
        await service.add_document("notes.pdf", "Volcano workshop notes.")

        with pytest.raises(ValueError, match="Upload or select"):
            await service.answer("What is in the notes?", [])

    asyncio.run(run())
